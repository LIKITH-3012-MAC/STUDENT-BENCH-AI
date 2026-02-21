import os
import io
import csv
import re
import zipfile
import tarfile
import gzip
import mimetypes
import openpyxl
from PIL import Image
import pytesseract
from docx import Document
from flask import Flask, request, jsonify, render_template, session
from groq import Groq
from dotenv import load_dotenv
from PyPDF2 import PdfReader
from flask_cors import CORS
from flask_session import Session

# ================= INIT =================
load_dotenv()

app = Flask(__name__)

# ================= CONFIG =================
app.secret_key = os.getenv("SECRET_KEY", "prometheus_super_secret_key")

app.config["SESSION_TYPE"] = "filesystem"
app.config["SESSION_PERMANENT"] = False
app.config["SESSION_USE_SIGNER"] = True
app.config["SESSION_FILE_DIR"] = "./.flask_session/"
app.config["SESSION_FILE_THRESHOLD"] = 100

Session(app)
CORS(app, supports_credentials=True)

# ================= GROQ =================
client = Groq(api_key=os.getenv("GROQ_API_KEY"))
MODEL_NAME = os.getenv("MODEL_NAME", "llama-3.1-8b-instant")

# ================= CONSTANTS =================
CREATOR_NAME = "Likith Naidu Anumakonda"
BOT_NAME = "Prometheus AI"
IDENTITY_RESPONSE = f"I was developed and created by {CREATOR_NAME}."

MAX_FILE_SIZE_MB = 25
MAX_ARCHIVE_FILES = 20

# ================= SECURITY =================
def is_identity_question(message: str) -> bool:
    message = message.lower()
    patterns = [
        r"who (developed|created|built|made) you",
        r"who is your (creator|owner|god)",
        r"who owns you",
        r"your developer",
        r"your creator",
        r"who made you",
    ]
    return any(re.search(pattern, message) for pattern in patterns)

# ================= SYSTEM PROMPT =================
BASE_SYSTEM_PROMPT = f"""
You are {BOT_NAME}, an intelligent AI assistant.

You were developed and created by {CREATOR_NAME}.

Always remain helpful, professional, and intelligent.
Never mention Groq, Meta, OpenAI, LLaMA, or backend providers.
"""

# ================= HOME =================
@app.route("/")
def home():
    return render_template("index.html")

# ================= HEALTH =================
@app.route("/health")
def health():
    return jsonify({
        "status": "OK",
        "bot": BOT_NAME,
        "creator": CREATOR_NAME,
        "model": MODEL_NAME
    }), 200

# ================= BINARY CHECK =================
def is_binary_string(bytes_data):
    textchars = bytearray({7, 8, 9, 10, 12, 13, 27} | set(range(0x20, 0x100)))
    return bool(bytes_data.translate(None, textchars))

# ================= UNIVERSAL FILE EXTRACTOR =================
def extract_text_from_file(file):
    filename = file.filename.lower()

    # File size protection
    file.seek(0, os.SEEK_END)
    file_size = file.tell()
    file.seek(0)

    if file_size > MAX_FILE_SIZE_MB * 1024 * 1024:
        return "⚠️ File too large."

    # ===== PDF =====
    if filename.endswith(".pdf"):
        reader = PdfReader(file)
        return "\n".join(page.extract_text() or "" for page in reader.pages)

    # ===== DOCX =====
    if filename.endswith(".docx"):
        doc = Document(file)
        return "\n".join(p.text for p in doc.paragraphs)

    # ===== EXCEL =====
    if filename.endswith((".xlsx", ".xlsm", ".xltx", ".xltm")):
        wb = openpyxl.load_workbook(file)
        text = ""
        for sheet in wb.worksheets:
            text += f"\n--- Sheet: {sheet.title} ---\n"
            for row in sheet.iter_rows(values_only=True):
                text += " ".join(str(cell) for cell in row if cell) + "\n"
        return text

    # ===== CSV =====
    if filename.endswith(".csv"):
        stream = io.StringIO(file.read().decode("utf-8", errors="ignore"))
        reader = csv.reader(stream)
        return "\n".join(" ".join(row) for row in reader)

    # ===== IMAGES (OCR) =====
    if filename.endswith((".png", ".jpg", ".jpeg", ".bmp", ".tiff", ".webp")):
        image = Image.open(file)
        return pytesseract.image_to_string(image)

    # ===== ZIP =====
    if filename.endswith(".zip"):
        text = ""
        with zipfile.ZipFile(file) as z:
            for i, name in enumerate(z.namelist()):
                if i >= MAX_ARCHIVE_FILES:
                    break
                with z.open(name) as f:
                    content = f.read()
                    if not is_binary_string(content):
                        text += f"\n--- {name} ---\n"
                        text += content.decode("utf-8", errors="ignore")
        return text

    # ===== TAR =====
    if filename.endswith((".tar", ".tar.gz", ".tgz")):
        text = ""
        with tarfile.open(fileobj=file) as tar:
            for i, member in enumerate(tar.getmembers()):
                if i >= MAX_ARCHIVE_FILES:
                    break
                f = tar.extractfile(member)
                if f:
                    content = f.read()
                    if not is_binary_string(content):
                        text += f"\n--- {member.name} ---\n"
                        text += content.decode("utf-8", errors="ignore")
        return text

    # ===== GZIP =====
    if filename.endswith(".gz"):
        with gzip.open(file, "rb") as f:
            content = f.read()
            if not is_binary_string(content):
                return content.decode("utf-8", errors="ignore")

    # ===== GENERIC TEXT =====
    try:
        raw_bytes = file.read()
        file.seek(0)
        if not is_binary_string(raw_bytes):
            return raw_bytes.decode("utf-8", errors="ignore")
    except:
        pass

    return "⚠️ Unsupported or binary file format."

# ================= CHAT =================
@app.route("/chat", methods=["POST"])
def chat():
    try:
        data = request.get_json()
        user_message = data.get("message", "").strip()

        if not user_message:
            return jsonify({"reply": "⚠️ Empty message"}), 400

        if is_identity_question(user_message):
            return jsonify({"reply": IDENTITY_RESPONSE})

        context = session.get("file_text", "")

        messages = [{"role": "system", "content": BASE_SYSTEM_PROMPT}]

        if context:
            messages.append({
                "role": "system",
                "content": f"Use this document if relevant:\n{context[:12000]}"
            })

        messages.append({"role": "user", "content": user_message})

        response = client.chat.completions.create(
            model=MODEL_NAME,
            messages=messages,
            temperature=0.5,
            max_tokens=1000
        )

        return jsonify({"reply": response.choices[0].message.content})

    except Exception as e:
        return jsonify({"reply": f"⚠️ Error: {str(e)}"}), 500

# ================= UPLOAD =================
@app.route("/upload", methods=["POST"])
def upload():
    try:
        if "file" not in request.files:
            return jsonify({"reply": "⚠️ No file provided"}), 400

        file = request.files["file"]
        user_query = request.form.get("query", "Summarize this document.")

        if not file.filename:
            return jsonify({"reply": "⚠️ No file selected"}), 400

        extracted_text = extract_text_from_file(file)

        if not extracted_text or not extracted_text.strip():
            return jsonify({"reply": "⚠️ Could not extract text."}), 400

        session["file_text"] = extracted_text

        if is_identity_question(user_query):
            return jsonify({"reply": IDENTITY_RESPONSE})

        messages = [
            {"role": "system", "content": BASE_SYSTEM_PROMPT},
            {
                "role": "user",
                "content": f"Document:\n{extracted_text[:12000]}\n\nQuestion: {user_query}"
            }
        ]

        response = client.chat.completions.create(
            model=MODEL_NAME,
            messages=messages,
            temperature=0.4
        )

        return jsonify({
            "reply": response.choices[0].message.content,
            "status": "success"
        })

    except Exception as e:
        return jsonify({"reply": f"⚠️ Upload error: {str(e)}"}), 500

# ================= CLEAR MEMORY =================
@app.route("/clear", methods=["POST"])
def clear():
    session.pop("file_text", None)
    return jsonify({"reply": "Memory cleared!"})

# ================= RUN =================
if __name__ == "__main__":
    app.run(host="0.0.0.0", port=5000)
